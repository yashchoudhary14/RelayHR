"""Create a disposable one-slide fixture for P0 tool verification."""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt


def main():
    destination = Path(__file__).resolve().parents[1] / "tmp" / "p0_smoke.pptx"
    destination.parent.mkdir(parents=True, exist_ok=True)
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(11), Inches(2))
    shape.name = "smoke_title"
    run = shape.text_frame.paragraphs[0].add_run()
    run.text = "RelayHR workspace render check"
    run.font.size = Pt(32)
    presentation.save(destination)
    print(destination)


if __name__ == "__main__":
    main()
