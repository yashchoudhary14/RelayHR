"""Deterministic QA for the submission deck.

Checks slide count, shapes that leave the slide, text below a minimum font size,
text boxes without an explicit font size, words per slide, and leftover placeholders.

Usage: python tools/check_pptx.py <deck.pptx> [--content-slides 2] [--min-pt 10] [--max-words 230]
Exit code 1 when any check fails.
"""
import argparse
import re
import sys

from pptx import Presentation
from pptx.util import Emu

PLACEHOLDER = re.compile(r"\{[A-Z_0-9]+\}|TODO|TBD|XXX|lorem", re.IGNORECASE)


def iter_shapes(shapes):
    for shp in shapes:
        yield shp
        if shp.shape_type == 6:  # group
            yield from iter_shapes(shp.shapes)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx")
    ap.add_argument("--content-slides", type=int, default=2,
                    help="content slides allowed, excluding cover and thank-you")
    ap.add_argument("--min-pt", type=float, default=10.0,
                    help="smallest allowed body font size")
    ap.add_argument("--footnote-pt", type=float, default=8.0,
                    help="smallest allowed size for shapes named 'footnote*' or 'source*'")
    ap.add_argument("--max-words", type=int, default=230)
    ap.add_argument("--allow-placeholder", action="append", default=[], metavar="NAME",
                    help="tolerate {NAME} while its value is not known yet (e.g. DOC_URL); repeatable")
    args = ap.parse_args()
    allowed = {f"{{{name.strip('{}').upper()}}}" for name in args.allow_placeholder}

    prs = Presentation(args.pptx)
    sw, sh = prs.slide_width, prs.slide_height
    failures = []
    n = len(prs.slides)
    print(f"Slide size: {Emu(sw).inches:.2f} x {Emu(sh).inches:.2f} in; slides: {n}")
    if n > args.content_slides + 2:
        failures.append(f"{n} slides > cover + {args.content_slides} content + thank-you")

    for idx, slide in enumerate(prs.slides, start=1):
        words = 0
        for shp in iter_shapes(slide.shapes):
            name = shp.name or ""
            if shp.left is not None and shp.width is not None:
                if shp.left < 0 or shp.top < 0 or shp.left + shp.width > sw or shp.top + shp.height > sh:
                    failures.append(f"slide {idx}: '{name}' extends beyond the slide edge")
            if not shp.has_text_frame:
                continue
            floor = args.footnote_pt if name.lower().startswith(("footnote", "source")) else args.min_pt
            for para in shp.text_frame.paragraphs:
                text = "".join(r.text for r in para.runs).strip()
                if not text:
                    continue
                words += len(text.split())
                if any(m.group(0) not in allowed for m in PLACEHOLDER.finditer(text)):
                    failures.append(f"slide {idx}: placeholder text in '{name}': {text[:60]!r}")
                for run in para.runs:
                    if not run.text.strip():
                        continue
                    size = run.font.size
                    if size is None:
                        failures.append(f"slide {idx}: '{name}' has a run without explicit font size: {run.text[:40]!r}")
                    elif size.pt < floor:
                        failures.append(f"slide {idx}: '{name}' uses {size.pt}pt (< {floor}pt): {run.text[:40]!r}")
        print(f"slide {idx}: {words} words")
        if 1 < idx < n and words > args.max_words:
            failures.append(f"slide {idx}: {words} words > {args.max_words}")
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            print(f"  notes: {len(notes.split())} words")

    if failures:
        print("\nFAIL")
        for f in failures:
            print(" -", f)
        sys.exit(1)
    print("\nPASS")


if __name__ == "__main__":
    main()
